"""
Отчёт «Проблемки» — пропадания связи аптеки по данным Zabbix
(триггер из настроек, по умолчанию [NET] NO PING IN VPN). В БД ничего не пишется.
"""
import io
from datetime import date, datetime, timedelta

from django.http import HttpResponse
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from ..models import Client, SystemSettings
from ..zabbix import ZabbixError, get_outages, format_duration, local_tz

DEFAULT_DAYS = 30
MAX_DAYS = 366


def _parse_period(request, settings):
    """date_from/date_to (YYYY-MM-DD) → даты; по умолчанию последние 30 дней."""
    today = datetime.now(local_tz(settings)).date()
    try:
        date_to = date.fromisoformat(request.query_params.get('date_to') or today.isoformat())
        date_from = date.fromisoformat(
            request.query_params.get('date_from') or (date_to - timedelta(days=DEFAULT_DAYS - 1)).isoformat()
        )
    except ValueError:
        raise ZabbixError('Неверный формат даты (ожидается ГГГГ-ММ-ДД)')
    if date_from > date_to:
        raise ZabbixError('Дата начала позже даты окончания')
    if (date_to - date_from).days + 1 > MAX_DAYS:
        raise ZabbixError(f'Период не может быть больше {MAX_DAYS} дней')
    return date_from, date_to


def _build_report(request, pk):
    try:
        client = Client.objects.get(pk=pk)
    except Client.DoesNotExist:
        raise ZabbixError('Клиент не найден')
    settings = SystemSettings.get()
    date_from, date_to = _parse_period(request, settings)
    report = get_outages(settings, client.pharmacy_code, date_from, date_to)
    return client, report


class ClientProblemsView(APIView):
    """GET /api/clients/{pk}/problems/?date_from=&date_to= — пропадания связи за период."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk=None):
        try:
            _, report = _build_report(request, pk)
        except ZabbixError as e:
            return Response({'error': str(e)}, status=400)
        return Response(report)


class ClientProblemsExportView(APIView):
    """GET /api/clients/{pk}/problems/export/?date_from=&date_to= — тот же отчёт в Excel."""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk=None):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        try:
            client, report = _build_report(request, pk)
        except ZabbixError as e:
            return Response({'error': str(e)}, status=400)

        fmt = lambda iso: datetime.fromisoformat(iso).strftime('%d.%m.%Y %H:%M')  # noqa: E731
        period = f"{fmt(report['period']['from'])} — {fmt(report['period']['to'])}"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Проблемки'

        hfont = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        hfill = PatternFill('solid', start_color='1677FF')
        halign = Alignment(horizontal='center', vertical='center', wrap_text=True)
        calign = Alignment(vertical='center')
        thin = Side(style='thin', color='CCCCCC')
        brd = Border(left=thin, right=thin, top=thin, bottom=thin)
        font = Font(name='Arial', size=10)
        bold = Font(name='Arial', size=10, bold=True)

        # Шапка отчёта
        info = [
            ('Аптека', f"{report['host']} — {client.address or report['host_name']}"),
            ('Триггер', report['trigger']),
            ('Период', period),
            ('Связь сейчас', 'НЕТ' if report['current_problem'] else 'есть'),
        ]
        for ri, (label, value) in enumerate(info, 1):
            ws.cell(row=ri, column=1, value=label).font = bold
            ws.cell(row=ri, column=2, value=value).font = font

        headers = [('№', 6), ('Пропал', 20), ('Появился', 20), ('Длительность', 18), ('Минут', 10)]
        head_row = len(info) + 2
        for col, (title, width) in enumerate(headers, 1):
            cell = ws.cell(row=head_row, column=col, value=title)
            cell.font = hfont; cell.fill = hfill; cell.alignment = halign; cell.border = brd
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width, 45)

        row = head_row
        for i, it in enumerate(report['items'], 1):
            row = head_row + i
            start = fmt(it['start']) + (' (до периода)' if it['started_before'] else '')
            end = 'нет связи сейчас' if it['open'] else fmt(it['end'])
            values = [i, start, end, format_duration(it['duration_sec']), round(it['duration_sec'] / 60)]
            fill = PatternFill('solid', start_color='F8FBFF' if i % 2 == 0 else 'FFFFFF')
            for col, v in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.font = font; cell.alignment = calign; cell.border = brd; cell.fill = fill

        total_row = row + 2
        ws.cell(row=total_row, column=1, value='Итого').font = bold
        ws.cell(row=total_row, column=2,
                value=f"{report['total_count']} пропаданий, простой {format_duration(report['total_sec'])} "
                      f"({report['percent']}% времени)").font = bold

        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

        buf = io.BytesIO()
        wb.save(buf)
        d_from = report['period']['from'][:10]
        d_to = report['period']['to'][:10]
        filename = f"problemki_{report['host']}_{d_from}_{d_to}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
